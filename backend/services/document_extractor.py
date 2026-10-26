import os
import fitz
from docx import Document as WordDocument
from pptx import Presentation


def extract_pdf_pages(file_path):
    """Extract text from a PDF page by page."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    pages = []
    document = fitz.open(file_path)

    try:
        for page_index in range(len(document)):
            page = document[page_index]
            pages.append({
                "page_number": page_index + 1,
                "text": page.get_text("text").strip()
            })
    finally:
        document.close()

    return pages


def extract_document_pages(file_path, extension):
    """Extract text from PDF, DOCX, PPTX, or TXT."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    extension = extension.lower().lstrip(".")

    if extension == "pdf":
        return extract_pdf_pages(file_path)

    if extension == "docx":
        document = WordDocument(file_path)
        text_parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())

        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells
                )
                if row_text.strip(" |"):
                    text_parts.append(row_text)

        return [{"page_number": 1, "text": "\n".join(text_parts)}]

    if extension == "pptx":
        presentation = Presentation(file_path)
        pages = []

        for slide_number, slide in enumerate(presentation.slides, start=1):
            text_parts = []

            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    text_parts.append(shape.text.strip())

                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(
                            cell.text.strip() for cell in row.cells
                        )
                        if row_text.strip(" |"):
                            text_parts.append(row_text)

            pages.append({
                "page_number": slide_number,
                "text": "\n".join(text_parts)
            })

        return pages

    if extension == "txt":
        with open(file_path, "r", encoding="utf-8-sig") as file:
            text = file.read().strip()

        return [{"page_number": 1, "text": text}]

    raise ValueError(f"Unsupported document format: {extension}")
